import io
import json
import unittest
from unittest.mock import patch

from langchain_core.documents import Document
from rag.evidence import EvidenceUnit, validate_grounded_answer


class AnswerCorrectnessTests(unittest.TestCase):
    def check(self, evidence, answer):
        unit = EvidenceUnit(evidence, {'semantic_relevance': 0.95})
        return validate_grounded_answer('นโยบายสินค้า A', answer, [evidence], [unit])[0]

    def test_negated_policy_cannot_be_reversed(self):
        self.assertFalse(self.check('นโยบายคืนสินค้า: ไม่รับคืนสินค้าทุกกรณี', 'รับคืนสินค้าทุกกรณี'))

    def test_prices_cannot_be_reassigned(self):
        self.assertFalse(self.check('ราคา: สินค้า A 100 บาท สินค้า B 200 บาท', 'สินค้า A ราคา 200 บาท'))

    def test_invented_assertion_without_numbers_is_rejected(self):
        self.assertFalse(self.check('นโยบาย: รับคืนสินค้าที่ชำรุด', 'จัดส่งฟรีทั่วประเทศ'))

    def test_condition_cannot_be_removed(self):
        self.assertFalse(self.check('เฉพาะสินค้าที่ชำรุด รับคืนสินค้าได้', 'รับคืนสินค้าได้'))

    def test_empty_evidence_is_rejected(self):
        self.assertFalse(validate_grounded_answer('นโยบาย', 'รับคืนได้', [], [])[0])

    def test_exact_policy_remains_accepted(self):
        self.assertTrue(self.check('นโยบายคืนสินค้า: ไม่รับคืนสินค้าทุกกรณี', 'ไม่รับคืนสินค้าทุกกรณี'))

    def test_quote_cannot_crop_out_negation_from_source(self):
        unit = EvidenceUnit('นโยบาย: ไม่รับคืนสินค้า', {'semantic_relevance': 0.95})
        self.assertFalse(validate_grounded_answer('นโยบาย', 'รับคืนสินค้า', ['รับคืนสินค้า'], [unit])[0])

    def test_second_negation_does_not_mask_deleted_first_negation(self):
        self.assertFalse(self.check('ไม่รับคืนสินค้าที่ไม่มีใบเสร็จ', 'รับคืนสินค้าที่ไม่มีใบเสร็จ'))

    def test_condition_on_previous_line_is_preserved(self):
        self.assertFalse(self.check('เฉพาะสินค้าชำรุด\nรับคืนสินค้าได้', 'รับคืนสินค้าได้'))

    def test_thai_leading_vowel_cannot_reverse_open_to_close(self):
        self.assertFalse(self.check('เปิดร้านทุกวัน', 'ปิดร้านทุกวัน'))

    def test_english_negation_is_preserved(self):
        self.assertFalse(self.check('Returns are not accepted', 'accepted'))


class WordTableTests(unittest.TestCase):
    def test_table_values_and_headers_survive_in_document_order(self):
        import docx
        from rag.document_loader import load_docx
        document = docx.Document()
        document.add_paragraph('ก่อนตาราง')
        table = document.add_table(rows=3, cols=2)
        for row, values in zip(table.rows, [('สินค้า', 'ราคา'), ('A', '100 บาท'), ('B', '200 บาท')]):
            for cell, value in zip(row.cells, values):
                cell.text = value
        document.add_paragraph('หลังตาราง')
        stream = io.BytesIO()
        document.save(stream)
        stream.seek(0)
        text = load_docx(stream)
        self.assertIn('สินค้า: A', text)
        self.assertIn('ราคา: 100 บาท', text)
        self.assertLess(text.index('ก่อนตาราง'), text.index('100 บาท'))
        self.assertLess(text.index('200 บาท'), text.index('หลังตาราง'))
        from rag.chunker import split_documents
        chunks = split_documents([Document(page_content=text, metadata={'source': 'table.docx'})])
        row = next(c.page_content for c in chunks if '100 บาท' in c.page_content)
        self.assertIn('สินค้า: A', row)
        self.assertNotIn('200 บาท', row)

    def test_headerless_and_nested_tables_do_not_lose_text(self):
        import docx
        from rag.document_loader import load_docx
        document = docx.Document()
        table = document.add_table(rows=2, cols=2)
        table.cell(0, 0).text = 'A'
        table.cell(0, 1).text = '100 บาท'
        table.cell(1, 0).text = 'B'
        table.cell(1, 1).text = '200 บาท'
        nested = table.cell(1, 1).add_table(rows=1, cols=1)
        nested.cell(0, 0).text = 'รวมภาษี'
        stream = io.BytesIO()
        document.save(stream)
        stream.seek(0)
        text = load_docx(stream)
        for expected in ('A', '100 บาท', 'B', '200 บาท', 'รวมภาษี'):
            self.assertIn(expected, text)
        self.assertNotIn('100 บาท:', text)


class GroundingPipelineTests(unittest.TestCase):
    def test_model_negation_reversal_is_not_published(self):
        from rag import rag_pipeline
        doc = Document(page_content='นโยบายคืนสินค้า: ไม่รับคืนสินค้าทุกกรณี',
                       metadata={'source': 'policy.txt', 'semantic_relevance': 0.95})
        with patch.object(rag_pipeline, 'retrieve_docs', return_value=[doc]), \
             patch.object(rag_pipeline.ollama, 'chat', return_value={'message': {'content': json.dumps({
                 'answer': 'รับคืนสินค้าทุกกรณี', 'evidence_ids': [1]})}}):
            result = rag_pipeline.ask_rag('รับคืนสินค้าได้ไหม', 'test')
        self.assertEqual(result, ('REQUIRE_HUMAN_HANDOFF', []))


class CompleteListTests(unittest.TestCase):
    def units(self, count):
        from rag.evidence import split_evidence_units
        return split_evidence_units([Document(page_content='นโยบายคืนสินค้า\n' + '\n'.join(
            '- รับคืนสินค้าเงื่อนไขที่ ' + str(i) for i in range(1, count + 1)), metadata={'source': 'policy.txt'})])

    def test_unnumbered_section_keeps_all_twelve_items(self):
        from rag.evidence import select_evidence_units, extract_section_answer
        units = self.units(12)
        selected = select_evidence_units('นโยบายคืนสินค้ามีอะไรบ้าง', units)
        self.assertEqual(len(selected), 12)
        answer = extract_section_answer('นโยบายคืนสินค้ามีอะไรบ้าง', selected)
        self.assertIsNotNone(answer)
        self.assertEqual(len(answer.answer.splitlines()), 12)

    def test_over_limit_is_marked_incomplete_not_silently_cut(self):
        from rag.evidence import select_evidence_units
        selected = select_evidence_units('นโยบายคืนสินค้ามีอะไรบ้าง', self.units(81))
        self.assertTrue(any(u.metadata.get('section_incomplete') for u in selected))

    def test_unnumbered_section_survives_second_pass(self):
        from rag.evidence import split_evidence_units
        units = self.units(12)
        second = split_evidence_units([Document(page_content=u.text, metadata=u.metadata) for u in units])
        self.assertEqual(units, second)

    def test_retrieval_expands_one_hit_into_complete_unnumbered_section(self):
        from rag.retriever import retrieve_docs
        from rag.semantic import QueryPlan
        from rag import rag_pipeline
        from unittest.mock import MagicMock
        question = 'นโยบายคืนสินค้ามีอะไรบ้าง'
        units = self.units(12)
        docs = [Document(page_content=u.text, metadata=u.metadata) for u in units]
        store = MagicMock()
        store.similarity_search_with_score.return_value = [(docs[0], 0.1)]
        store.get_section_documents.return_value = docs
        with patch('api.knowledge.resolve_index_path', return_value='unused'), \
             patch('rag.retriever.os.path.isfile', return_value=True), \
             patch('rag.retriever.get_vector_db', return_value=store), \
             patch('rag.semantic.prepare_query', return_value=QueryPlan(question, (question,))), \
             patch('rag.semantic.rank_candidates', return_value=[docs[0]]):
            retrieved = retrieve_docs(question, 'test')
        self.assertEqual(len(retrieved), 12)
        with patch.object(rag_pipeline, 'retrieve_docs', return_value=retrieved), \
             patch.object(rag_pipeline.ollama, 'chat') as model:
            answer, sources = rag_pipeline.ask_rag(question, 'test')
        self.assertEqual(len(answer.splitlines()), 12)
        self.assertTrue(sources)
        model.assert_not_called()


class FollowupContextTests(unittest.TestCase):
    def test_short_warranty_followup_uses_subject_for_ranking(self):
        from rag.semantic import prepare_query
        with patch('rag.semantic.rewrite_client') as client:
            client.return_value.chat.return_value = {'message': {'content': json.dumps({
                'question': 'สินค้า A รับประกันกี่ปี', 'queries': ['สินค้า A รับประกัน']})}}
            plan = prepare_query('รับประกันกี่ปี', [{'question': 'สินค้า A ราคาเท่าไร'}])
        self.assertIn('สินค้า A', plan.question)
        self.assertIn('รับประกันกี่ปี', plan.question)
        self.assertIn(plan.question, plan.queries)

    def test_rewrite_outage_keeps_followup_context(self):
        from rag.semantic import prepare_query
        with patch('rag.semantic.rewrite_client', side_effect=RuntimeError('offline')):
            plan = prepare_query('รับประกันกี่ปี', [{'question': 'สินค้า A ราคาเท่าไร', 'answer': 'SECRET'}])
        self.assertIn('สินค้า A', plan.question)
        self.assertNotIn('SECRET', plan.question)
        self.assertIn(plan.question, plan.queries)

    def test_explicit_new_product_does_not_inherit_previous_one(self):
        from rag.semantic import prepare_query
        with patch('rag.semantic.rewrite_client', side_effect=RuntimeError('offline')):
            plan = prepare_query('สินค้า B รับประกันกี่ปี', [{'question': 'สินค้า A ราคาเท่าไร'}])
        self.assertEqual(plan.question, 'สินค้า B รับประกันกี่ปี')

    def test_consecutive_followups_keep_latest_explicit_subject(self):
        from rag.semantic import prepare_query
        with patch('rag.semantic.rewrite_client', side_effect=RuntimeError('offline')):
            plan = prepare_query('รับประกันกี่ปี', [{'question': 'สินค้า A มีอะไรบ้าง'}, {'question': 'ราคาเท่าไร'}])
        self.assertIn('สินค้า A', plan.question)


if __name__ == '__main__':
    unittest.main()
